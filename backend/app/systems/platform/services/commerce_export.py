"""订单收款、退款记录导出为 Excel。"""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_SHANGHAI = ZoneInfo("Asia/Shanghai")

ORDER_STATUS_LABELS = {
    "pending": "待支付",
    "paid": "已收款",
    "refunded": "已退款",
    "cancelled": "已取消",
}
ORDER_TYPE_LABELS = {
    "membership": "会籍办卡",
    "retail": "零售",
    "pt": "私教",
    "pt_package": "私教课包",
    "group": "团课",
    "dining": "餐饮消费",
    "coupon": "优惠券",
    "course_pack": "课程包",
    "activity": "活动报名",
}
REFUND_STATUS_LABELS = {
    "succeeded": "已到账",
    "processing": "处理中",
    "created": "已发起",
    "failed": "失败",
}
REFUND_CHANNEL_LABELS = {
    "wechat_original": "退回微信",
    "offline_cash": "现金",
    "offline_transfer": "转账",
    "online": "线上",
}


def shanghai_text(value: datetime | None) -> str:
    if value is None:
        return ""
    if value.tzinfo is None:
        value = value.replace(tzinfo=ZoneInfo("UTC"))
    return value.astimezone(_SHANGHAI).strftime("%Y-%m-%d %H:%M:%S")


def label_of(mapping: dict[str, str], code: str | None) -> str:
    if not code:
        return ""
    return mapping.get(code, code)


def build_table_xlsx(sheet_title: str, headers: list[str], rows: list[list]) -> bytes:
    """生成单表 xlsx：表头深色，金额列保持数字。"""
    wb = Workbook()
    sheet = wb.active
    sheet.title = sheet_title[:31]
    header_fill = PatternFill("solid", fgColor="171B1F")
    header_font = Font(bold=True, color="F2E6D2")
    for col, title in enumerate(headers, start=1):
        cell = sheet.cell(1, col, title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[get_column_letter(col)].width = 18 if len(title) <= 4 else 22
    for row_index, row in enumerate(rows, start=2):
        for col, value in enumerate(row, start=1):
            cell = sheet.cell(row_index, col, value)
            if isinstance(value, Decimal):
                cell.value = float(value)
                cell.number_format = "0.00"
            cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
